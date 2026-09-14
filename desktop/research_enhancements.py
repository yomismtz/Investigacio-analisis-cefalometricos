from __future__ import annotations

import json
import sqlite3
import tempfile
import zipfile
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
import tkinter as tk

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from research_db import RESEARCH_DIR, ResearchDB, now_iso
from research_protocol import selected_variables

_INSTALLED = False


def _refresh_case_status_exact(self: ResearchDB, case_id: int, _selected_measurements=None) -> str:
    """Mark complete only when every exact analysis::measurement key has a value."""
    row = self.case(case_id)
    if not row:
        raise KeyError(case_id)
    if not row["included"]:
        return "excluded"
    protocol = self.latest_protocol(row["study_id"])
    required = set((protocol.get("config") or {}).get("selected_variables") or [])
    results = self.latest_results(case_id, protocol["version"])
    valid = {
        f"{r['analysis']}::{r['measurement']}"
        for r in results
        if r["value"] is not None
    }
    if required and required.issubset(valid):
        status = "complete"
    elif results or self.latest_landmarks(case_id, protocol["version"]):
        status = "incomplete"
    else:
        status = "pending"
    with self.connect() as con:
        con.execute(
            "UPDATE cases SET status=?,updated_at=? WHERE id=?",
            (status, now_iso(), case_id),
        )
    return status


def _case_criteria_values(self: ResearchDB, case_id: int) -> dict[int, bool]:
    with self.connect() as con:
        rows = con.execute(
            "SELECT criterion_id,value FROM case_criteria WHERE case_id=?", (case_id,)
        ).fetchall()
    return {int(r["criterion_id"]): bool(r["value"]) for r in rows}


def _backup_study_consistent(self: ResearchDB, study_id: int, target_zip: str | Path) -> str:
    """Create an SQLite-consistent snapshot even while WAL/autosave is active."""
    target_zip = str(target_zip)
    study_folder = RESEARCH_DIR / f"study_{study_id:04d}"
    with tempfile.TemporaryDirectory(prefix="yomceph_backup_") as td:
        snapshot = Path(td) / "yomceph_research.sqlite3"
        src = sqlite3.connect(self.db_path)
        dst = sqlite3.connect(snapshot)
        try:
            src.backup(dst)
            dst.commit()
        finally:
            dst.close()
            src.close()
        with zipfile.ZipFile(target_zip, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(snapshot, arcname="database/yomceph_research.sqlite3")
            if study_folder.exists():
                for p in study_folder.rglob("*"):
                    if p.is_file():
                        z.write(p, arcname=str(Path("study_files") / p.relative_to(study_folder)))
    self.audit(study_id, "backup_created", target=target_zip, sqlite_snapshot=True)
    return target_zip


def export_excel_with_landmarks(db: ResearchDB, study_id: int, protocol: dict, path: str) -> str:
    """Workbook ready for statistics plus long-form landmark coordinates."""
    wb = Workbook()
    violet = "6D35A4"
    mint = "45C7A5"
    header_fill = PatternFill("solid", fgColor=violet)
    mint_fill = PatternFill("solid", fgColor=mint)

    ws = wb.active
    ws.title = "Datos"
    variables = selected_variables(protocol.get("selected_variables") or [])
    headers = [
        "case_number", "study_code", "sex", "age", "status", "included",
        "exclusion_reason", "examiner",
    ] + [v.key for v in variables]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    version = protocol.get("version", db.latest_protocol(study_id)["version"])
    for c in db.list_cases(study_id, include_excluded=True):
        values = {
            f"{r['analysis']}::{r['measurement']}": r
            for r in db.latest_results(c["id"], version)
        }
        row = [
            c["case_number"], c["study_code"], c["sex"], c["age"], c["status"],
            c["included"], c["exclusion_reason"], c["examiner"],
        ]
        row += [None if values.get(v.key) is None else values[v.key]["value"] for v in variables]
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(36, len(h) + 2))

    lws = wb.create_sheet("Landmarks")
    lheaders = ["case_number", "study_code", "status", "landmark", "x_px", "y_px", "protocol_version"]
    lws.append(lheaders)
    for cell in lws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = mint_fill
        cell.alignment = Alignment(horizontal="center")
    for c in db.list_cases(study_id, include_excluded=True):
        pts = db.latest_landmarks(c["id"], version)
        for name, (x, y) in pts.items():
            lws.append([c["case_number"], c["study_code"], c["status"], name, x, y, version])
    lws.freeze_panes = "A2"
    lws.auto_filter.ref = lws.dimensions
    for i, width in enumerate([14, 18, 14, 24, 16, 16, 18], 1):
        lws.column_dimensions[get_column_letter(i)].width = width

    pws = wb.create_sheet("Protocolo")
    pws.append(["Campo", "Valor"])
    for cell in pws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for k, v in protocol.items():
        pws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
    pws.column_dimensions["A"].width = 30
    pws.column_dimensions["B"].width = 100

    aws = wb.create_sheet("Auditoría")
    aws.append(["id", "case_id", "evento", "detalles", "fecha"])
    for cell in aws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for r in db.audit_rows(study_id):
        aws.append([r["id"], r["case_id"], r["event"], r["details_json"], r["created_at"]])
    aws.column_dimensions["D"].width = 80

    wb.save(path)
    db.audit(study_id, "export_excel", target=path, landmarks_sheet=True)
    return path


def _case_checklist_dialog(self):
    cid = self.selected_case_id() if hasattr(self, "selected_case_id") else None
    if not cid:
        messagebox.showinfo("YomCeph", self.T("Selecciona un caso primero.", "Select a case first."), parent=self)
        return
    criteria = list(self.db.criteria(self.study_id))
    existing = self.db.case_criteria_values(cid)
    w = tk.Toplevel(self)
    w.title(self.T("Checklist del caso", "Case checklist"))
    w.transient(self)
    w.grab_set()
    w.geometry("650x620")
    ttk.Label(
        w,
        text=self.T("Criterios de inclusión y exclusión", "Inclusion and exclusion criteria"),
        font=("Segoe UI", 14, "bold"),
    ).pack(anchor="w", padx=14, pady=(14, 6))
    ttk.Label(
        w,
        text=self.T(
            "Marca únicamente los criterios que realmente cumple este caso. El registro queda en auditoría.",
            "Check only criteria that actually apply to this case. The record is retained in audit.",
        ),
        wraplength=610,
    ).pack(anchor="w", padx=14, pady=(0, 8))
    canvas = tk.Canvas(w, highlightthickness=0)
    scroll = ttk.Scrollbar(w, orient="vertical", command=canvas.yview)
    inner = ttk.Frame(canvas)
    win = canvas.create_window((0, 0), window=inner, anchor="nw")
    canvas.configure(yscrollcommand=scroll.set)
    canvas.pack(side="left", fill="both", expand=True, padx=(14, 0), pady=(0, 60))
    scroll.pack(side="right", fill="y", pady=(0, 60))
    inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.bind("<Configure>", lambda e: canvas.itemconfigure(win, width=e.width))
    vars_ = {}
    current_kind = None
    for r in criteria:
        if r["kind"] != current_kind:
            current_kind = r["kind"]
            ttk.Label(
                inner,
                text=self.T("INCLUSIÓN" if current_kind == "include" else "EXCLUSIÓN", "INCLUSION" if current_kind == "include" else "EXCLUSION"),
                font=("Segoe UI", 10, "bold"),
            ).pack(anchor="w", pady=(10, 3))
        v = tk.BooleanVar(value=existing.get(int(r["id"]), False))
        vars_[int(r["id"])] = v
        ttk.Checkbutton(inner, text=r["label"], variable=v).pack(anchor="w", fill="x", pady=2)

    footer = ttk.Frame(w)
    footer.place(relx=0, rely=1, relwidth=1, anchor="sw", height=55)
    def save():
        values = {cid_: var.get() for cid_, var in vars_.items()}
        self.db.set_case_criteria(cid, values)
        w.destroy()
        self.status.config(text=self.T("Checklist del caso guardado", "Case checklist saved"))
    ttk.Button(footer, text=self.T("Guardar checklist", "Save checklist"), command=save).pack(side="right", padx=14, pady=10)
    ttk.Button(footer, text=self.T("Cancelar", "Cancel"), command=w.destroy).pack(side="right", padx=4, pady=10)


def _export_xlsx_enhanced(self):
    if not self.study_id:
        return
    p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
    if p:
        export_excel_with_landmarks(self.db, self.study_id, self.protocol_for_export(), p)
        messagebox.showinfo("YomCeph", self.T("Excel completo exportado, incluyendo landmarks X/Y.", "Complete Excel exported, including X/Y landmarks."), parent=self)


def install(workspace_class=None):
    global _INSTALLED
    if _INSTALLED:
        return
    ResearchDB.refresh_case_status = _refresh_case_status_exact
    ResearchDB.case_criteria_values = _case_criteria_values
    ResearchDB.backup_study = _backup_study_consistent
    if workspace_class is not None:
        workspace_class.case_checklist_dialog = _case_checklist_dialog
        workspace_class.export_xlsx = _export_xlsx_enhanced
        old_build_cases = workspace_class._build_cases
        def build_cases_with_checklist(self):
            old_build_cases(self)
            ttk.Button(
                self.cases_tab,
                text=self.T("Checklist de inclusión/exclusión del caso seleccionado", "Inclusion/exclusion checklist for selected case"),
                command=self.case_checklist_dialog,
            ).pack(fill="x", pady=(6, 0))
        workspace_class._build_cases = build_cases_with_checklist
        old_bindings = workspace_class._bindings
        def bindings_with_checklist(self):
            old_bindings(self)
            self.bind("<Control-k>", lambda e: self.case_checklist_dialog())
        workspace_class._bindings = bindings_with_checklist
    _INSTALLED = True
