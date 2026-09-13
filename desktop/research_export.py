from __future__ import annotations

import csv
import json
from collections import OrderedDict
from pathlib import Path

from research_protocol import selected_variables


def _latest_value_map(db, case_id: int, protocol_version: int) -> dict[str, dict]:
    rows = db.latest_results(case_id, protocol_version)
    return {f"{r['analysis']}::{r['measurement']}": dict(r) for r in rows}


def export_wide_csv(db, study_id: int, protocol: dict, path: str, include_excluded: bool = False) -> str:
    keys = list(protocol.get("selected_variables") or [])
    variables = selected_variables(keys)
    cases = db.list_cases(study_id, include_excluded=include_excluded)
    headers = ["case_number", "study_code", "sex", "age", "status", "included", "exclusion_reason", "examiner"] + [v.key for v in variables]
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=headers)
        w.writeheader()
        for c in cases:
            values = _latest_value_map(db, c["id"], protocol.get("version", db.latest_protocol(study_id)["version"]))
            row = {
                "case_number": c["case_number"], "study_code": c["study_code"], "sex": c["sex"], "age": c["age"],
                "status": c["status"], "included": c["included"], "exclusion_reason": c["exclusion_reason"], "examiner": c["examiner"],
            }
            for v in variables:
                item = values.get(v.key)
                row[v.key] = "" if item is None or item.get("value") is None else item.get("value")
            w.writerow(row)
    db.audit(study_id, "export_wide_csv", target=path, include_excluded=include_excluded)
    return path


def export_long_csv(db, study_id: int, protocol: dict, path: str, include_excluded: bool = False) -> str:
    cases = db.list_cases(study_id, include_excluded=include_excluded)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        fields = ["case_number", "study_code", "sex", "age", "status", "included", "analysis", "measurement", "value", "unit", "reference", "interpretation", "examiner"]
        w = csv.DictWriter(fh, fieldnames=fields); w.writeheader()
        version = protocol.get("version", db.latest_protocol(study_id)["version"])
        for c in cases:
            for r in db.latest_results(c["id"], version):
                w.writerow({
                    "case_number": c["case_number"], "study_code": c["study_code"], "sex": c["sex"], "age": c["age"], "status": c["status"], "included": c["included"],
                    "analysis": r["analysis"], "measurement": r["measurement"], "value": r["value"], "unit": r["unit"], "reference": r["reference_text"], "interpretation": r["interpretation"], "examiner": c["examiner"],
                })
    db.audit(study_id, "export_long_csv", target=path, include_excluded=include_excluded)
    return path


def export_audit_csv(db, study_id: int, path: str) -> str:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh); w.writerow(["id", "study_id", "case_id", "event", "details_json", "created_at"])
        for r in db.audit_rows(study_id):
            w.writerow([r["id"], r["study_id"], r["case_id"], r["event"], r["details_json"], r["created_at"]])
    return path


def export_json_bundle(db, study_id: int, path: str) -> str:
    study = dict(db.study(study_id))
    protocol = db.latest_protocol(study_id)
    cases = []
    for c in db.list_cases(study_id, include_excluded=True):
        item = dict(c)
        item["landmarks"] = db.latest_landmarks(c["id"], protocol["version"])
        item["results"] = [dict(r) for r in db.latest_results(c["id"], protocol["version"])]
        cases.append(item)
    data = {"format": "YomCeph Research Bundle", "version": 1, "study": study, "protocol": protocol, "cases": cases, "audit": [dict(r) for r in db.audit_rows(study_id)]}
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    db.audit(study_id, "export_json_bundle", target=path)
    return path


def export_excel(db, study_id: int, protocol: dict, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = "Datos"
    variables = selected_variables(protocol.get("selected_variables") or [])
    headers = ["case_number", "study_code", "sex", "age", "status", "included", "exclusion_reason", "examiner"] + [v.key for v in variables]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="6D35A4")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True); cell.fill = header_fill; cell.alignment = Alignment(horizontal="center")
    version = protocol.get("version", db.latest_protocol(study_id)["version"])
    for c in db.list_cases(study_id, include_excluded=True):
        values = _latest_value_map(db, c["id"], version)
        base = [c["case_number"], c["study_code"], c["sex"], c["age"], c["status"], c["included"], c["exclusion_reason"], c["examiner"]]
        base += [None if values.get(v.key) is None else values[v.key]["value"] for v in variables]
        ws.append(base)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(34, len(h) + 2))

    pws = wb.create_sheet("Protocolo")
    pws.append(["Campo", "Valor"])
    pws["A1"].font = pws["B1"].font = Font(bold=True, color="FFFFFF")
    pws["A1"].fill = pws["B1"].fill = header_fill
    for k, v in protocol.items():
        pws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
    aws = wb.create_sheet("Auditoría")
    aws.append(["id", "case_id", "evento", "detalles", "fecha"])
    for r in db.audit_rows(study_id): aws.append([r["id"], r["case_id"], r["event"], r["details_json"], r["created_at"]])
    wb.save(path)
    db.audit(study_id, "export_excel", target=path)
    return path


def export_case_pdf(db, case_id: int, protocol: dict, path: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    case = db.case(case_id)
    study = db.study(case["study_id"])
    version = protocol.get("version", db.latest_protocol(case["study_id"])["version"])
    rows = db.latest_results(case_id, version)
    styles = getSampleStyleSheet(); doc = SimpleDocTemplate(path, pagesize=A4)
    story = [Paragraph("YomCeph Desktop — Informe de caso", styles["Title"]), Spacer(1, 8)]
    story.append(Paragraph(f"Estudio: {study['title']} · Caso: {case['study_code']} · Estado: {case['status']}", styles["BodyText"]))
    story.append(Paragraph("Uso educativo e investigación. No es un dispositivo médico ni sustituye el diagnóstico profesional.", styles["Italic"]))
    data = [["Análisis", "Medición", "Valor", "Unidad", "Referencia"]]
    for r in rows: data.append([r["analysis"], r["measurement"], "" if r["value"] is None else f"{r['value']:.3f}", r["unit"], r["reference_text"]])
    t = Table(data, repeatRows=1, colWidths=[90, 145, 60, 45, 175]); t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#6D35A4")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#CDB7EE")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("FONTSIZE", (0,0), (-1,-1), 7)])); story += [Spacer(1, 10), t]
    doc.build(story); db.audit(case["study_id"], "export_case_pdf", case_id=case_id, target=path); return path


def export_study_pdf(db, study_id: int, protocol: dict, path: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    study = db.study(study_id); progress = db.progress(study_id); styles = getSampleStyleSheet(); doc = SimpleDocTemplate(path, pagesize=A4)
    story = [Paragraph("YomCeph Desktop — Resumen de investigación", styles["Title"]), Paragraph(study["title"], styles["Heading2"]), Spacer(1, 8)]
    story.append(Paragraph(f"Investigador: {study['researcher']} · Institución: {study['institution']}", styles["BodyText"]))
    story.append(Paragraph("Uso educativo e investigación. No es un dispositivo médico ni sustituye el diagnóstico profesional.", styles["Italic"]))
    summary = [["Total", "Completos", "Pendientes", "En análisis", "Incompletos", "Excluidos"], [progress['total'], progress['complete'], progress['pending'], progress['in_progress'], progress['incomplete'], progress['excluded']]]
    t = Table(summary); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#6D35A4")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.25,colors.HexColor("#CDB7EE")),("ALIGN",(0,0),(-1,-1),"CENTER")]))
    story += [Spacer(1,10), t, Spacer(1,12), Paragraph("Variables seleccionadas", styles["Heading3"])]
    for v in selected_variables(protocol.get("selected_variables") or []): story.append(Paragraph(f"• {v.analysis}: {v.name}", styles["BodyText"]))
    doc.build(story); db.audit(study_id, "export_study_pdf", target=path); return path
