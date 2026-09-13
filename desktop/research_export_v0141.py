from __future__ import annotations
import csv,json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def _case_version(db,c):
    if 'protocol_version' in c.keys() and c['protocol_version']: return int(c['protocol_version'])
    return db.case_protocol_version(int(c['id']))

def _protocols_and_keys(db,study_id):
    protocols=db.protocol_versions(study_id); keys=[]
    for p in protocols:
        for k in (p.get('config') or {}).get('selected_variables') or []:
            if k not in keys: keys.append(k)
    return protocols,keys

def export_wide(db,study_id,path,include_excluded=False):
    _,keys=_protocols_and_keys(db,study_id)
    headers=['case_number','study_code','sex','age','protocol_version','status','included','exclusion_reason','examiner']+keys
    with open(path,'w',newline='',encoding='utf-8-sig') as fh:
        w=csv.DictWriter(fh,fieldnames=headers);w.writeheader()
        for c in db.list_cases(study_id,include_excluded=include_excluded):
            pv=_case_version(db,c); vals={f"{r['analysis']}::{r['measurement']}":r['value'] for r in db.latest_results(c['id'],pv)}
            row={'case_number':c['case_number'],'study_code':c['study_code'],'sex':c['sex'],'age':c['age'],'protocol_version':pv,'status':c['status'],'included':c['included'],'exclusion_reason':c['exclusion_reason'],'examiner':c['examiner']}
            row.update({k:'' if vals.get(k) is None else vals[k] for k in keys});w.writerow(row)
    db.audit(study_id,'export_wide_csv',target=path,include_excluded=include_excluded,version_aware=True);return path

def export_long(db,study_id,path,include_excluded=False):
    fields=['case_number','study_code','sex','age','protocol_version','status','included','analysis','measurement','value','unit','reference','interpretation','examiner']
    with open(path,'w',newline='',encoding='utf-8-sig') as fh:
        w=csv.DictWriter(fh,fieldnames=fields);w.writeheader()
        for c in db.list_cases(study_id,include_excluded=include_excluded):
            pv=_case_version(db,c)
            for r in db.latest_results(c['id'],pv):
                w.writerow({'case_number':c['case_number'],'study_code':c['study_code'],'sex':c['sex'],'age':c['age'],'protocol_version':pv,'status':c['status'],'included':c['included'],'analysis':r['analysis'],'measurement':r['measurement'],'value':r['value'],'unit':r['unit'],'reference':r['reference_text'],'interpretation':r['interpretation'],'examiner':c['examiner']})
    db.audit(study_id,'export_long_csv',target=path,include_excluded=include_excluded,version_aware=True);return path

def export_excel(db,study_id,path):
    protocols,keys=_protocols_and_keys(db,study_id);wb=Workbook();ws=wb.active;ws.title='Datos';violet='6D35A4';mint='45C7A5';fill=PatternFill('solid',fgColor=violet)
    headers=['case_number','study_code','sex','age','protocol_version','status','included','exclusion_reason','examiner']+keys;ws.append(headers)
    for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=fill;cell.alignment=Alignment(horizontal='center')
    for c in db.list_cases(study_id,include_excluded=True):
        pv=_case_version(db,c);vals={f"{r['analysis']}::{r['measurement']}":r['value'] for r in db.latest_results(c['id'],pv)}
        ws.append([c['case_number'],c['study_code'],c['sex'],c['age'],pv,c['status'],c['included'],c['exclusion_reason'],c['examiner']]+[vals.get(k) for k in keys])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=max(12,min(38,len(h)+2))
    lws=wb.create_sheet('Landmarks');lws.append(['case_number','study_code','status','protocol_version','landmark','x_px','y_px'])
    for cell in lws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor=mint)
    for c in db.list_cases(study_id,include_excluded=True):
        pv=_case_version(db,c)
        for name,(x,y) in db.latest_landmarks(c['id'],pv).items():lws.append([c['case_number'],c['study_code'],c['status'],pv,name,x,y])
    pws=wb.create_sheet('Protocolos');pws.append(['version','locked','created_at','config_json'])
    for cell in pws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=fill
    for p in protocols:pws.append([p['version'],int(p['locked']),p['created_at'],json.dumps(p['config'],ensure_ascii=False)])
    pws.column_dimensions['D'].width=100
    aws=wb.create_sheet('Auditoría');aws.append(['id','case_id','evento','detalles','fecha'])
    for cell in aws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=fill
    for r in db.audit_rows(study_id):aws.append([r['id'],r['case_id'],r['event'],r['details_json'],r['created_at']])
    aws.column_dimensions['D'].width=80;wb.save(path);db.audit(study_id,'export_excel',target=path,version_aware=True,landmarks_sheet=True);return path

def export_json(db,study_id,path):
    cases=[]
    for c in db.list_cases(study_id,include_excluded=True):
        pv=_case_version(db,c);item=dict(c);item['protocol_version']=pv;item['landmarks']=db.latest_landmarks(c['id'],pv);item['results']=[dict(r) for r in db.latest_results(c['id'],pv)];cases.append(item)
    data={'format':'YomCeph Research Bundle','version':2,'study':dict(db.study(study_id)),'protocols':db.protocol_versions(study_id),'cases':cases,'audit':[dict(r) for r in db.audit_rows(study_id)]}
    Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8');db.audit(study_id,'export_json_bundle',target=path,version_aware=True);return path

def export_case_pdf(db,case_id,path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    case=db.case(case_id);study=db.study(case['study_id']);pv=db.case_protocol_version(case_id);rows=db.latest_results(case_id,pv);styles=getSampleStyleSheet();doc=SimpleDocTemplate(path,pagesize=A4)
    story=[Paragraph('YomCeph Desktop — Informe de caso',styles['Title']),Spacer(1,8),Paragraph(f"Estudio: {study['title']} · Caso: {case['study_code']} · Protocolo v{pv} · Estado: {case['status']}",styles['BodyText']),Paragraph('Uso educativo e investigación. No es un dispositivo médico ni sustituye el diagnóstico profesional.',styles['Italic'])]
    data=[['Análisis','Medición','Valor','Unidad','Referencia']]
    for r in rows:data.append([r['analysis'],r['measurement'],'' if r['value'] is None else f"{r['value']:.3f}",r['unit'],r['reference_text']])
    t=Table(data,repeatRows=1,colWidths=[90,145,60,45,175]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#6D35A4')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#CDB7EE')),('VALIGN',(0,0),(-1,-1),'TOP'),('FONTSIZE',(0,0),(-1,-1),7)]));story += [Spacer(1,10),t];doc.build(story);db.audit(case['study_id'],'export_case_pdf',case_id=case_id,target=path,protocol_version=pv);return path
